"""Fill scraped Rakuten Travel rates into the "03_料金採取テンプレ" sheet of the
survey workbook (沖縄県恩納村エリア_競合料金調査_再構築フレームワーク.xlsx).

Input records (one per scraped room plan, typically the JSON emitted by
`python -m scraper.cli --mode hotel`) look like:

    {"hotel_name": "ハレクラニ沖縄", "week": "W1", "plan_name": "...",
     "price": "35,000円", "room_type": "T2"}

`room_type` is optional; if missing it is suggested from `plan_name` via
`roomtype.classify_room_type`. Records with no resolvable room_type, or whose
hotel_name doesn't match a row in the sheet, are reported as warnings and
skipped rather than guessed at - the 05 sheet asks operators to hand-check
ambiguous rooms rather than force a bucket.

If two records land on the same cell (e.g. a weekday + weekend sample for the
same hotel/week/room-type, per the 05 sheet's sampling rule), their prices are
averaged.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass

import openpyxl

from .roomtype import classify_room_type

SHEET_NAME = "03_料金採取テンプレ"
HOTEL_NAME_COL = 2  # column B
FIRST_HOTEL_ROW = 7
LAST_HOTEL_ROW = 26

WEEK_ROOMTYPE_COLS = {
    "W1": {"T1": 3, "T2": 4, "T3": 5, "T4": 6},
    "W2": {"T1": 7, "T2": 8, "T3": 9, "T4": 10},
    "W3": {"T1": 11, "T2": 12, "T3": 13, "T4": 14},
}


@dataclass
class FillWarning:
    record: dict
    reason: str


def _normalize(name: str) -> str:
    return re.sub(r"\s+", "", name or "").lower()


def _parse_price(price) -> float | None:
    if isinstance(price, (int, float)):
        return float(price)
    if not price:
        return None
    digits = re.sub(r"[^\d]", "", str(price))
    return float(digits) if digits else None


def find_hotel_row(ws, hotel_name: str) -> int | None:
    target = _normalize(hotel_name)
    if not target:
        return None
    for row in range(FIRST_HOTEL_ROW, LAST_HOTEL_ROW + 1):
        cell_name = ws.cell(row=row, column=HOTEL_NAME_COL).value
        if cell_name and _normalize(cell_name) == target:
            return row
    for row in range(FIRST_HOTEL_ROW, LAST_HOTEL_ROW + 1):
        cell_name = ws.cell(row=row, column=HOTEL_NAME_COL).value
        if cell_name and (target in _normalize(cell_name) or _normalize(cell_name) in target):
            return row
    return None


def fill_from_records(ws, records: list[dict]) -> tuple[int, list[FillWarning]]:
    """Write `records` into `ws`, averaging values that land on the same cell.

    Returns (number of cells written, list of skipped-record warnings).
    """
    cell_values: dict[tuple[int, int], list[float]] = {}
    warnings: list[FillWarning] = []

    for record in records:
        hotel_name = record.get("hotel_name")
        week = record.get("week")
        room_type = record.get("room_type") or classify_room_type(record.get("plan_name", ""))
        price = _parse_price(record.get("price"))

        if not hotel_name:
            warnings.append(FillWarning(record, "hotel_name is missing"))
            continue
        if week not in WEEK_ROOMTYPE_COLS:
            warnings.append(FillWarning(record, f"week must be one of W1/W2/W3, got {week!r}"))
            continue
        if not room_type or room_type not in WEEK_ROOMTYPE_COLS[week]:
            warnings.append(FillWarning(record, f"could not resolve a T1-T4 room_type for plan {record.get('plan_name')!r}"))
            continue
        if price is None:
            warnings.append(FillWarning(record, f"could not parse a numeric price from {record.get('price')!r}"))
            continue

        row = find_hotel_row(ws, hotel_name)
        if row is None:
            warnings.append(FillWarning(record, f"no matching hotel row for {hotel_name!r}"))
            continue

        col = WEEK_ROOMTYPE_COLS[week][room_type]
        cell_values.setdefault((row, col), []).append(price)

    for (row, col), values in cell_values.items():
        ws.cell(row=row, column=col).value = sum(values) / len(values)

    return len(cell_values), warnings


def fill_template(xlsx_path: str, data_paths: list[str], output_path: str) -> tuple[int, list[FillWarning]]:
    records: list[dict] = []
    for path in data_paths:
        with open(path, encoding="utf-8") as f:
            payload = json.load(f)
        records.extend(payload if isinstance(payload, list) else [payload])

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    filled, warnings = fill_from_records(ws, records)
    wb.save(output_path)
    return filled, warnings


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fill scraped rates into the survey workbook's 03 sheet.")
    parser.add_argument("--xlsx", required=True, help="Path to the survey workbook (.xlsx)")
    parser.add_argument("--data", required=True, nargs="+", help="One or more JSON result files (from scraper.cli --mode hotel)")
    parser.add_argument("--output", required=True, help="Path to write the filled workbook to")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    filled, warnings = fill_template(args.xlsx, args.data, args.output)
    print(f"Filled {filled} cell(s) in {SHEET_NAME} -> {args.output}")
    for w in warnings:
        print(f"WARNING: {w.reason}: {w.record}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
