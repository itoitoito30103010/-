"""Fill the 03_料金採取テンプレ sheet of the competitive-rate Excel workbook."""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# Column mapping in sheet 03_料金採取テンプレ (1-indexed, matching openpyxl convention)
# Row 5 = header row with T1/T2/T3/T4 labels
# Row 6 = date labels
# Rows 7-26 = hotel data (20 hotels)
# Columns: A=#, B=name, C=W1/T1, D=W1/T2, E=W1/T3, F=W1/T4,
#           G=W2/T1, H=W2/T2, I=W2/T3, J=W2/T4,
#           K=W3/T1, L=W3/T2, M=W3/T3, N=W3/T4, O=URL

WEEK_TYPE_TO_COL = {
    ("W1", "T1"): 3,
    ("W1", "T2"): 4,
    ("W1", "T3"): 5,
    ("W1", "T4"): 6,
    ("W2", "T1"): 7,
    ("W2", "T2"): 8,
    ("W2", "T3"): 9,
    ("W2", "T4"): 10,
    ("W3", "T1"): 11,
    ("W3", "T2"): 12,
    ("W3", "T3"): 13,
    ("W3", "T4"): 14,
}

DATE_ROW = 6
DATA_START_ROW = 7
HOTEL_NAME_COL = 2
URL_COL = 15


def fill_template(
    template_path: str | Path,
    output_path: str | Path,
    results: list[dict],
    weeks: dict[str, tuple[str, str]],
) -> None:
    """
    Write scraped rates into the Excel template and save to output_path.

    results: list of dicts as returned by scraper.scrape_hotel_rates()
    weeks:   {"W1": ("2026-08-12", "2026-08-13"), ...}
    """
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required: pip install openpyxl")

    wb = openpyxl.load_workbook(template_path)
    ws = wb["03_料金採取テンプレ"]

    # Write week dates into row 6
    week_date_cols = {"W1": 3, "W2": 7, "W3": 11}
    for week_label, col in week_date_cols.items():
        if week_label in weeks:
            checkin, checkout = weeks[week_label]
            ws.cell(row=DATE_ROW, column=col).value = f"{checkin}〜{checkout}"

    # Build a lookup from hotel name → result row
    name_to_result = {r["name"]: r for r in results}

    for row_idx in range(DATA_START_ROW, DATA_START_ROW + 20):
        hotel_name_cell = ws.cell(row=row_idx, column=HOTEL_NAME_COL).value
        if not hotel_name_cell:
            continue
        hotel_name = str(hotel_name_cell).strip()
        row_data = name_to_result.get(hotel_name)
        if row_data is None:
            logger.warning("No scraped data for hotel %r (row %d)", hotel_name, row_idx)
            continue

        for (week_label, room_type), col in WEEK_TYPE_TO_COL.items():
            week_rates = row_data.get(week_label, {})
            price = week_rates.get(room_type)
            if price is not None:
                ws.cell(row=row_idx, column=col).value = price

    wb.save(output_path)
    logger.info("Saved filled workbook to %s", output_path)
