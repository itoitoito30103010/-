import unittest

import openpyxl

from scraper.fill_template import fill_from_records, find_hotel_row

HOTELS = [
    "ハレクラニ沖縄",
    "ハイアットリージェンシー瀬良垣アイランド沖縄",
    "カフーリゾートフチャクコンド・ホテル",
]


def _make_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "03_料金採取テンプレ"
    for i, name in enumerate(HOTELS):
        ws.cell(row=7 + i, column=1).value = i + 1
        ws.cell(row=7 + i, column=2).value = name
    return wb, ws


class FindHotelRowTest(unittest.TestCase):
    def test_exact_match(self):
        _, ws = _make_sheet()
        self.assertEqual(find_hotel_row(ws, "ハレクラニ沖縄"), 7)

    def test_fuzzy_match_ignores_whitespace(self):
        _, ws = _make_sheet()
        self.assertEqual(find_hotel_row(ws, "ハイアット リージェンシー瀬良垣アイランド沖縄"), 8)

    def test_no_match_returns_none(self):
        _, ws = _make_sheet()
        self.assertIsNone(find_hotel_row(ws, "存在しないホテル"))


class FillFromRecordsTest(unittest.TestCase):
    def test_writes_price_to_correct_cell(self):
        _, ws = _make_sheet()
        records = [
            {"hotel_name": "ハレクラニ沖縄", "week": "W1", "room_type": "T2", "price": "50,000円"},
        ]
        filled, warnings = fill_from_records(ws, records)
        self.assertEqual(filled, 1)
        self.assertEqual(warnings, [])
        self.assertEqual(ws.cell(row=7, column=4).value, 50000.0)

    def test_averages_multiple_samples_on_same_cell(self):
        _, ws = _make_sheet()
        records = [
            {"hotel_name": "ハレクラニ沖縄", "week": "W1", "room_type": "T1", "price": "10000"},
            {"hotel_name": "ハレクラニ沖縄", "week": "W1", "room_type": "T1", "price": "20000"},
        ]
        filled, warnings = fill_from_records(ws, records)
        self.assertEqual(filled, 1)
        self.assertEqual(warnings, [])
        self.assertEqual(ws.cell(row=7, column=3).value, 15000.0)

    def test_suggests_room_type_from_plan_name_when_missing(self):
        _, ws = _make_sheet()
        records = [
            {"hotel_name": "カフーリゾートフチャクコンド・ホテル", "week": "W3", "plan_name": "ヴィラタイプ", "price": "8000"},
        ]
        filled, warnings = fill_from_records(ws, records)
        self.assertEqual(filled, 1)
        self.assertEqual(warnings, [])
        self.assertEqual(ws.cell(row=9, column=14).value, 8000.0)

    def test_unresolvable_room_type_is_a_warning_not_a_guess(self):
        _, ws = _make_sheet()
        records = [
            {"hotel_name": "ハレクラニ沖縄", "week": "W1", "plan_name": "特別プラン", "price": "8000"},
        ]
        filled, warnings = fill_from_records(ws, records)
        self.assertEqual(filled, 0)
        self.assertEqual(len(warnings), 1)
        self.assertIn("room_type", warnings[0].reason)

    def test_unknown_hotel_is_a_warning(self):
        _, ws = _make_sheet()
        records = [
            {"hotel_name": "存在しないホテル", "week": "W1", "room_type": "T1", "price": "8000"},
        ]
        filled, warnings = fill_from_records(ws, records)
        self.assertEqual(filled, 0)
        self.assertEqual(len(warnings), 1)
        self.assertIn("no matching hotel row", warnings[0].reason)

    def test_unparseable_price_is_a_warning(self):
        _, ws = _make_sheet()
        records = [
            {"hotel_name": "ハレクラニ沖縄", "week": "W1", "room_type": "T1", "price": "満室"},
        ]
        filled, warnings = fill_from_records(ws, records)
        self.assertEqual(filled, 0)
        self.assertEqual(len(warnings), 1)
        self.assertIn("parse", warnings[0].reason)


if __name__ == "__main__":
    unittest.main()
